import pandas as pd
from playwright.sync_api import sync_playwright
import time
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
import json
import urllib.parse
from datetime import datetime

# Import classifier and clean_text from existing script
from scraper_and_classifier import classify_content, clean_text, TAXONOMY

def scrape_article(page, url):
    try:
        page.goto(url, timeout=60000)
        time.sleep(3) # Wait for JS rendering
        
        # Extract title
        title = page.title()
        if " - " in title:
            title = title.split(" - ")[0]
        elif " | " in title:
            title = title.split(" | ")[0]
            
        # Extract content (paragraphs)
        paragraphs = page.locator("p").all_inner_texts()
        content = " ".join([p.strip() for p in paragraphs if p.strip()])
        
        # Extract publication
        pub = "Unknown"
        if "theprint.in" in url:
            pub = "ThePrint"
        elif "thehindu.com" in url:
            pub = "The Hindu"
        elif "hindustantimes.com" in url:
            pub = "Hindustan Times"
        elif "livemint.com" in url:
            pub = "Livemint"
        
        # Mock Date (extracting exact dates across different news layouts is flaky without specific selectors)
        date = datetime.now().strftime("%Y-%m-%d") 
        try:
            time_el = page.locator("time").first
            if time_el:
                datetime_str = time_el.get_attribute("datetime")
                if datetime_str:
                    date = datetime_str.split("T")[0]
        except:
            pass
            
        return {
            "title": clean_text(title),
            "content": clean_text(content),
            "publication": pub,
            "date": date
        }
    except Exception as e:
        print(f"Error scraping {url}: {e}")
        return {
            "title": "Failed to scrape",
            "content": f"Error: {e}",
            "publication": "Unknown",
            "date": datetime.now().strftime("%Y-%m-%d")
        }

def build_table_sheet(ws, dataset, tab_fill, headers, header_font, thin_border, zebra_fill, data_font, link_font, bold_data):
    ws.views.sheetView[0].showGridLines = True
    ws.append(headers)
    ws.row_dimensions[1].height = 28
    
    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.fill = tab_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
        
    for row_idx, item in enumerate(dataset, start=2):
        row_vals = [
            item["ID"], item["Date"], item["Publication"], item["Co-Authors"], item["Content Type"],
            item["Article Title"], item["Post Content"], item["URL"],
            item["Primary Classification"], item["Sub-classification"]
        ]
        ws.append(row_vals)
        ws.row_dimensions[row_idx].height = 65
        is_zebra = (row_idx % 2 == 0)
        
        for col_idx in range(1, len(headers) + 1):
            cell = ws.cell(row=row_idx, column=col_idx)
            cell.border = thin_border
            if is_zebra:
                cell.fill = zebra_fill
            
            col_name = headers[col_idx - 1]
            if col_name in ["ID", "Date", "Publication", "Co-Authors", "Content Type", "Primary Classification", "Sub-classification"]:
                cell.alignment = Alignment(horizontal="center", vertical="top")
                cell.font = data_font
            elif col_name in ["Article Title", "Post Content (Author Commentary)"]:
                cell.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
                cell.font = data_font
            elif col_name == "URL":
                cell.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
                cell.font = link_font
                cell.value = f'=HYPERLINK("{row_vals[7]}", "{row_vals[7]}")'
                
    widths = {
        "ID": 10, "Date": 13, "Publication": 20, "Co-Authors": 18, "Content Type": 18,
        "Article Title": 38, "Post Content (Author Commentary)": 50, "URL": 38,
        "Primary Classification": 24, "Sub-classification": 24
    }
    for col_idx, col_name in enumerate(headers, start=1):
        ws.column_dimensions[get_column_letter(col_idx)].width = widths[col_name]

def main():
    input_file = "manual/Untitled spreadsheet.xlsx"
    output_file = "manual/manual_master_sheet.xlsx"
    
    print(f"Reading {input_file}...")
    df_input = pd.read_excel(input_file)
    if 'shobhit sir article links' not in df_input.columns:
        print("Column 'shobhit sir article links' not found!")
        return
        
    links = df_input['shobhit sir article links'].dropna().tolist()
    print(f"Found {len(links)} links to process.")
    
    master_data = []
    
    with sync_playwright() as p:
        browser = p.chromium.launch(headless=True)
        # Using a context with user agent to bypass some bot checks
        context = browser.new_context(user_agent="Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/120.0.0.0 Safari/537.36")
        page = context.new_page()
        
        for i, url in enumerate(links):
            print(f"[{i+1}/{len(links)}] Scraping {url}...")
            article_data = scrape_article(page, url)
            
            # Content summary using first 500 chars for the sheet
            summary = article_data["content"][:500] + "..." if len(article_data["content"]) > 500 else article_data["content"]
            
            primary_cls, sub_cls = classify_content(article_data["content"], article_data["title"])
            
            master_data.append({
                "ID": f"MANUAL-{i+1:03d}",
                "Date": article_data["date"],
                "Publication": article_data["publication"],
                "Co-Authors": "Solo / Unknown",
                "Content Type": "Opinion Piece",
                "Article Title": article_data["title"],
                "Post Content": summary,
                "URL": url,
                "Primary Classification": primary_cls,
                "Sub-classification": sub_cls
            })
            
        browser.close()
        
    print("Scraping completed. Generating Master Sheet...")
    
    # Styles for Excel
    navy_fill = PatternFill(start_color="0F172A", end_color="0F172A", fill_type="solid")
    accent_fill = PatternFill(start_color="1E40AF", end_color="1E40AF", fill_type="solid")
    edu_fill = PatternFill(start_color="047857", end_color="047857", fill_type="solid")
    policy_fill = PatternFill(start_color="B45309", end_color="B45309", fill_type="solid")
    tech_fill = PatternFill(start_color="6D28D9", end_color="6D28D9", fill_type="solid")
    
    header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
    data_font = Font(name="Segoe UI", size=10)
    bold_data = Font(name="Segoe UI", size=10, bold=True)
    link_font = Font(name="Segoe UI", size=10, color="2563EB", underline="single")
    
    thin_border = Border(
        left=Side(style='thin', color='E2E8F0'),
        right=Side(style='thin', color='E2E8F0'),
        top=Side(style='thin', color='E2E8F0'),
        bottom=Side(style='thin', color='E2E8F0')
    )
    zebra_fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
    
    headers = [
        "ID", "Date", "Publication", "Co-Authors", "Content Type", 
        "Article Title", "Post Content (Author Commentary)", "URL", 
        "Primary Classification", "Sub-classification"
    ]
    
    wb = openpyxl.Workbook()
    
    # 1. Master Table Sheet
    ws_master = wb.active
    ws_master.title = "Master Database"
    build_table_sheet(ws_master, master_data, navy_fill, headers, header_font, thin_border, zebra_fill, data_font, link_font, bold_data)
    
    # 2. Category Tabs
    edu_data = [d for d in master_data if d["Primary Classification"] == "Education & Pedagogy"]
    ws_edu = wb.create_sheet(title="Education & Pedagogy")
    build_table_sheet(ws_edu, edu_data, edu_fill, headers, header_font, thin_border, zebra_fill, data_font, link_font, bold_data)
    
    policy_data = [d for d in master_data if d["Primary Classification"] == "Public Policy & Governance"]
    ws_policy = wb.create_sheet(title="Public Policy & Governance")
    build_table_sheet(ws_policy, policy_data, policy_fill, headers, header_font, thin_border, zebra_fill, data_font, link_font, bold_data)
    
    tech_data = [d for d in master_data if d["Primary Classification"] == "Technology & Innovation"]
    ws_tech = wb.create_sheet(title="Technology & Innovation")
    build_table_sheet(ws_tech, tech_data, tech_fill, headers, header_font, thin_border, zebra_fill, data_font, link_font, bold_data)
    
    wb.save(output_file)
    print(f"Successfully saved {len(master_data)} records to {output_file}")

if __name__ == "__main__":
    main()
