#!/usr/bin/env python3
"""
Convert CSV classifications to a professionally formatted Excel (.xlsx) workbook.
"""

import pandas as pd
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

def create_styled_excel():
    csv_file = "shobhit_mathur_article_classifications.csv"
    xlsx_file = "shobhit_mathur_article_classifications.xlsx"
    
    df = pd.read_csv(csv_file)
    
    wb = openpyxl.Workbook()
    
    # Define styles
    header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid") # Dark Navy Blue
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    data_font = Font(name="Calibri", size=10)
    link_font = Font(name="Calibri", size=10, color="0000EE", underline="single")
    bold_font = Font(name="Calibri", size=10, bold=True)
    
    thin_border = Border(
        left=Side(style='thin', color='E5E7EB'),
        right=Side(style='thin', color='E5E7EB'),
        top=Side(style='thin', color='E5E7EB'),
        bottom=Side(style='thin', color='E5E7EB')
    )
    
    zebra_fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
    
    # ----------------------------------------------------
    # Sheet 1: All Classified Articles
    # ----------------------------------------------------
    ws1 = wb.active
    ws1.title = "Classified Articles"
    ws1.views.sheetView[0].showGridLines = True
    
    headers = list(df.columns)
    ws1.append(headers)
    
    for col_idx in range(1, len(headers) + 1):
        cell = ws1.cell(row=1, column=col_idx)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws1.row_dimensions[1].height = 28
    
    for row_idx, row_data in enumerate(df.values, start=2):
        ws1.append(list(row_data))
        ws1.row_dimensions[row_idx].height = 65
        is_zebra = (row_idx % 2 == 0)
        
        for col_idx in range(1, len(headers) + 1):
            cell = ws1.cell(row=row_idx, column=col_idx)
            cell.border = thin_border
            if is_zebra:
                cell.fill = zebra_fill
            
            # Format columns
            col_name = headers[col_idx - 1]
            if col_name == "Date":
                cell.alignment = Alignment(horizontal="center", vertical="top")
                cell.font = data_font
            elif col_name in ["Post Content", "Article Title"]:
                cell.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
                cell.font = data_font
            elif col_name == "URL":
                cell.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
                cell.font = link_font
                cell.hyperlink = str(cell.value)
            elif col_name in ["Primary Classification", "Sub-classification"]:
                cell.alignment = Alignment(horizontal="center", vertical="top")
                cell.font = data_font
                
    # Column width configuration
    col_widths = {
        "Date": 14,
        "Post Content": 48,
        "Article Title": 36,
        "URL": 38,
        "Primary Classification": 24,
        "Sub-classification": 24
    }
    for col_idx, col_name in enumerate(headers, start=1):
        col_letter = get_column_letter(col_idx)
        ws1.column_dimensions[col_letter].width = col_widths.get(col_name, 20)

    # ----------------------------------------------------
    # Sheet 2: Article Links Quick Reference
    # ----------------------------------------------------
    ws2 = wb.create_sheet(title="Article Links Quick Ref")
    ws2.views.sheetView[0].showGridLines = True
    
    ref_headers = ["#", "Date", "Article Title", "Direct Link", "Primary Classification", "Sub-classification"]
    ws2.append(ref_headers)
    
    for col_idx in range(1, len(ref_headers) + 1):
        cell = ws2.cell(row=1, column=col_idx)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
    ws2.row_dimensions[1].height = 28
    
    for idx, row in df.iterrows():
        row_num = idx + 2
        ws2.append([
            idx + 1,
            row["Date"],
            row["Article Title"],
            row["URL"],
            row["Primary Classification"],
            row["Sub-classification"]
        ])
        ws2.row_dimensions[row_num].height = 24
        is_zebra = (row_num % 2 == 0)
        
        for col_idx in range(1, len(ref_headers) + 1):
            cell = ws2.cell(row=row_num, column=col_idx)
            cell.border = thin_border
            if is_zebra:
                cell.fill = zebra_fill
                
            if col_idx == 1: # #
                cell.alignment = Alignment(horizontal="center", vertical="center")
                cell.font = bold_font
            elif col_idx == 2: # Date
                cell.alignment = Alignment(horizontal="center", vertical="center")
                cell.font = data_font
            elif col_idx == 3: # Article Title
                cell.alignment = Alignment(horizontal="left", vertical="center")
                cell.font = data_font
            elif col_idx == 4: # Direct Link
                cell.alignment = Alignment(horizontal="left", vertical="center")
                cell.font = link_font
                cell.hyperlink = str(cell.value)
            else:
                cell.alignment = Alignment(horizontal="center", vertical="center")
                cell.font = data_font
                
    ws2.column_dimensions['A'].width = 6
    ws2.column_dimensions['B'].width = 14
    ws2.column_dimensions['C'].width = 50
    ws2.column_dimensions['D'].width = 55
    ws2.column_dimensions['E'].width = 25
    ws2.column_dimensions['F'].width = 25

    # ----------------------------------------------------
    # Sheet 3: Taxonomy & Analytics Summary
    # ----------------------------------------------------
    ws3 = wb.create_sheet(title="Category Summary")
    ws3.views.sheetView[0].showGridLines = True
    
    ws3.append(["Taxonomy Classification", "Count"])
    ws3.cell(row=1, column=1).fill = header_fill
    ws3.cell(row=1, column=1).font = header_font
    ws3.cell(row=1, column=2).fill = header_fill
    ws3.cell(row=1, column=2).font = header_font
    ws3.row_dimensions[1].height = 28
    
    summary_df = df["Primary Classification"].value_counts().reset_index()
    summary_df.columns = ["Primary Classification", "Count"]
    
    for idx, row in summary_df.iterrows():
        row_num = idx + 2
        ws3.append([row["Primary Classification"], row["Count"]])
        ws3.row_dimensions[row_num].height = 22
        ws3.cell(row=row_num, column=1).font = data_font
        ws3.cell(row=row_num, column=1).border = thin_border
        ws3.cell(row=row_num, column=2).font = bold_font
        ws3.cell(row=row_num, column=2).border = thin_border
        ws3.cell(row=row_num, column=2).alignment = Alignment(horizontal="center")
        
    total_row = len(summary_df) + 2
    ws3.append(["Total Articles", len(df)])
    ws3.row_dimensions[total_row].height = 24
    ws3.cell(row=total_row, column=1).font = bold_font
    ws3.cell(row=total_row, column=1).border = thin_border
    ws3.cell(row=total_row, column=2).font = bold_font
    ws3.cell(row=total_row, column=2).border = thin_border
    ws3.cell(row=total_row, column=2).alignment = Alignment(horizontal="center")
    
    ws3.column_dimensions['A'].width = 30
    ws3.column_dimensions['B'].width = 15
    
    # Save workbook
    wb.save(xlsx_file)
    print(f"Successfully generated styled Excel workbook: {xlsx_file}")

if __name__ == "__main__":
    create_styled_excel()
