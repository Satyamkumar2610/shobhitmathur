#!/usr/bin/env python3
"""
Generate a Master Sheet Excel workbook (shobhit_mathur_master_sheet.xlsx)
with unified metadata, dedicated category tabs, and an executive analytics dashboard.
"""

import pandas as pd
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# Extended master dataset with rich source and co-author metadata
MASTER_DATA = [
    {
        "ID": "SM-001",
        "Date": "2026-08-05",
        "Publication": "The Hindu",
        "Co-Authors": "Solo",
        "Content Type": "Opinion Column",
        "Article Title": "How the NEET can be restructured to end question paper leaks",
        "Post Content": "Examination leaks erode student morale and damage institutional credibility. Rebuilding faith in national entrance tests like NEET requires integrating modern computerized testing infrastructure with accountable state governance.",
        "URL": "https://www.thehindu.com/education/how-the-neet-can-be-restructured-to-end-question-paper-leaks/article71279765.ece",
        "Primary Classification": "Education & Pedagogy",
        "Sub-classification": "Systemic Reform"
    },
    {
        "ID": "SM-002",
        "Date": "2026-07-29",
        "Publication": "ThePrint",
        "Co-Authors": "Suresh Prabhu",
        "Content Type": "Opinion Column",
        "Article Title": "India topped the Physics Olympiad. Now it needs schools built for minds like these",
        "Post Content": "India's triumphant performance at the International Physics Olympiad demonstrates the immense raw intellect of our students. However, our standard classrooms and rote systems fail to nurture exceptional talent. We need specialized schools, talent incubators, and modern pedagogy designed specifically for extraordinary minds.",
        "URL": "https://theprint.in/opinion/india-topped-the-physics-olympiad-now-it-needs-schools-built-for-minds-like-these/",
        "Primary Classification": "Education & Pedagogy",
        "Sub-classification": "Pedagogy"
    },
    {
        "ID": "SM-003",
        "Date": "2026-07-07",
        "Publication": "ThePrint",
        "Co-Authors": "Suresh Prabhu",
        "Content Type": "Opinion Column",
        "Article Title": "What India can learn from America at 250",
        "Post Content": "As the United States marks 250 years, India should reflect on constitutional resilience, institutional balance, and state capacity. Building robust civic systems and transparent laws is foundational for sustainable nation-building.",
        "URL": "https://theprint.in/opinion/what-india-can-learn-from-america-at-250/",
        "Primary Classification": "Public Policy & Governance",
        "Sub-classification": "Nation-Building"
    },
    {
        "ID": "SM-004",
        "Date": "2026-06-19",
        "Publication": "ThePrint",
        "Co-Authors": "Suresh Prabhu",
        "Content Type": "Opinion Column",
        "Article Title": "India is mistaking data localisation for digital sovereignty. It must control the traffic system",
        "Post Content": "True digital sovereignty cannot be achieved merely by storing bits within national borders. We must master the digital traffic architecture, core algorithms, and communication layers to secure our tech rise and strategic independence.",
        "URL": "https://theprint.in/opinion/india-is-mistaking-data-localisation-for-digital-sovereignty-it-must-control-the-traffic-system/",
        "Primary Classification": "Technology & Innovation",
        "Sub-classification": "Digital Sovereignty"
    },
    {
        "ID": "SM-005",
        "Date": "2026-06-12",
        "Publication": "The Economic Times",
        "Co-Authors": "Suresh Prabhu",
        "Content Type": "Opinion Column",
        "Article Title": "India's development ladder after AI: What comes next?",
        "Post Content": "The traditional pathway of climbing the developmental ladder through low-cost services and manual labor is rapidly shifting due to artificial intelligence. How India navigates AI implications will determine our economic future.",
        "URL": "https://economictimes.indiatimes.com/opinion/et-commentary/indias-development-ladder-after-ai-what-comes-next/",
        "Primary Classification": "Technology & Innovation",
        "Sub-classification": "AI Implications"
    },
    {
        "ID": "SM-006",
        "Date": "2026-05-28",
        "Publication": "The Economic Times",
        "Co-Authors": "Suresh Prabhu",
        "Content Type": "Opinion Column",
        "Article Title": "India and the new economics of constraint",
        "Post Content": "In an era of fractured supply chains and shifting global alliances, India faces a new economics of constraint. We must rethink our regulatory frameworks and strategic domestic capacity to build economic resilience.",
        "URL": "https://economictimes.indiatimes.com/opinion/et-commentary/india-and-the-new-economics-of-constraint/",
        "Primary Classification": "Public Policy & Governance",
        "Sub-classification": "Regulatory Frameworks"
    },
    {
        "ID": "SM-007",
        "Date": "2026-04-21",
        "Publication": "ThePrint",
        "Co-Authors": "Suresh Prabhu",
        "Content Type": "Opinion Column",
        "Article Title": "University townships — India's chance to build its own version of what America got right",
        "Post Content": "Vibrant university towns serve as powerful engines of innovation, research, and local economic development. Developing integrated higher education campuses in India can transform regional growth.",
        "URL": "https://theprint.in/opinion/university-townships-indias-chance-to-build-its-own-version-of-what-america-got-right/",
        "Primary Classification": "Education & Pedagogy",
        "Sub-classification": "Higher Education"
    },
    {
        "ID": "SM-008",
        "Date": "2026-02-11",
        "Publication": "ThePrint",
        "Co-Authors": "Suresh Prabhu",
        "Content Type": "Opinion Column",
        "Article Title": "For-profit medical colleges are a step forward. Education is defined by outcomes, not ownership",
        "Post Content": "Expanding access to healthcare training requires enabling responsible private capital. Allowing for-profit medical colleges will bridge massive capacity shortages while regulatory frameworks ensure educational quality.",
        "URL": "https://theprint.in/opinion/for-profit-medical-colleges-are-a-step-forward-education-is-defined-by-outcomes-not-ownership/",
        "Primary Classification": "Education & Pedagogy",
        "Sub-classification": "Higher Education"
    },
    {
        "ID": "SM-009",
        "Date": "2025-06-18",
        "Publication": "The Economic Times",
        "Co-Authors": "Suresh Prabhu",
        "Content Type": "Opinion Column",
        "Article Title": "Technology education key to India's sovereignty",
        "Post Content": "Critical tech infrastructure demands homegrown talent and strategic educational programs. Fostering indigenous engineering and frugal innovation in our universities is essential to safeguard national digital security.",
        "URL": "https://economictimes.indiatimes.com/opinion/et-commentary/technology-education-key-to-indias-sovereignty/",
        "Primary Classification": "Technology & Innovation",
        "Sub-classification": "Frugal Innovation"
    },
    {
        "ID": "SM-010",
        "Date": "2025-06-10",
        "Publication": "The Economic Times",
        "Co-Authors": "Solo",
        "Content Type": "Opinion Column",
        "Article Title": "Fit to learn, fit to lead: Why India must rescue physical education now",
        "Post Content": "Declining physical engagement and sedentary academic environments hinder cognitive growth. Re-integrating sports and movement into our classrooms and curriculum is vital for student leadership.",
        "URL": "https://economictimes.indiatimes.com/opinion/et-commentary/fit-to-learn-fit-to-lead-why-india-must-rescue-physical-education-now/",
        "Primary Classification": "Education & Pedagogy",
        "Sub-classification": "Pedagogy"
    },
    {
        "ID": "SM-011",
        "Date": "2025-06-04",
        "Publication": "LinkedIn Pulse",
        "Co-Authors": "Solo",
        "Content Type": "Book Review",
        "Article Title": "Book Review | Yuk Hui’s The Question Concerning Technology in China (2016)",
        "Post Content": "What if technology is not universal—but civilizational? Yuk Hui’s groundbreaking book challenges one of modernity’s deepest assumptions: that technology is a neutral, universal force. A profound exploration of digital cosmotechnics and sovereignty.",
        "URL": "https://www.linkedin.com/pulse/book-review-yuk-huis-question-concerning-technology-china-shobhit-mathur/",
        "Primary Classification": "Technology & Innovation",
        "Sub-classification": "Digital Sovereignty"
    },
    {
        "ID": "SM-012",
        "Date": "2024-11-15",
        "Publication": "The Hindu",
        "Co-Authors": "Kankana Trivedi",
        "Content Type": "Research Essay",
        "Article Title": "Emerging trends in Public Policy Education in India",
        "Post Content": "India's rapid socioeconomic evolution demands specialized public policy education that bridges grassroots realities with rigorous institutional governance and evidence-based policymaking.",
        "URL": "https://www.thehindu.com/education/emerging-trends-in-public-policy-education-in-india/article68868921.ece",
        "Primary Classification": "Education & Pedagogy",
        "Sub-classification": "Higher Education"
    },
    {
        "ID": "SM-013",
        "Date": "2024-06-25",
        "Publication": "The Indian Express",
        "Co-Authors": "Suresh Prabhu",
        "Content Type": "Opinion Column",
        "Article Title": "To escape annual cycle of leaks and stress, India needs a new exam design",
        "Post Content": "High-stakes, single-day competitive exams create immense pressure on students and foster corrupt coaching cartels. We propose a computerized, multi-stage assessment system to drive genuine systemic reform in education policy.",
        "URL": "https://indianexpress.com/article/opinion/columns/to-escape-annual-cycle-of-leaks-and-stress-india-needs-a-new-exam-design/",
        "Primary Classification": "Education & Pedagogy",
        "Sub-classification": "Systemic Reform"
    },
    {
        "ID": "SM-014",
        "Date": "2024-06-19",
        "Publication": "Livemint",
        "Co-Authors": "Suresh Prabhu",
        "Content Type": "Opinion Column",
        "Article Title": "Educators should leverage AI for an educational transformation",
        "Post Content": "Artificial intelligence gives us an unprecedented opportunity to provide personalized tutoring at scale. Educational institutions and teachers must embrace EdTech and GenAI tools to personalize classroom pedagogy.",
        "URL": "https://www.livemint.com/opinion/online-views/educators-should-leverage-ai-for-an-educational-transformation-11718798993204.html",
        "Primary Classification": "Education & Pedagogy",
        "Sub-classification": "Higher Education"
    },
    {
        "ID": "SM-015",
        "Date": "2024-05-14",
        "Publication": "Hindustan Times",
        "Co-Authors": "Solo",
        "Content Type": "Opinion Column",
        "Article Title": "India needs policy support for entrepreneurial universities",
        "Post Content": "Transitioning Indian higher education from degree factories to entrepreneurship hubs requires responsive regulatory frameworks and government policies that support student startups and venture creation.",
        "URL": "https://www.hindustantimes.com/opinion/india-needs-policy-support-for-entrepreneurial-universities/",
        "Primary Classification": "Education & Pedagogy",
        "Sub-classification": "Higher Education"
    },
    {
        "ID": "SM-016",
        "Date": "2024-03-08",
        "Publication": "NDTV",
        "Co-Authors": "Solo",
        "Content Type": "Opinion Column",
        "Article Title": "Why India needs high-performing private universities with public purpose",
        "Post Content": "Private universities must step up to meet global standards while grounding education in civic responsibility and nation-building values. Regulatory bodies should incentivize institutional autonomy.",
        "URL": "https://www.ndtv.com/opinion/why-india-needs-high-performing-private-universities-with-public-purpose-5198234",
        "Primary Classification": "Education & Pedagogy",
        "Sub-classification": "Higher Education"
    },
    {
        "ID": "SM-017",
        "Date": "2023-11-20",
        "Publication": "Swarajya",
        "Co-Authors": "Solo",
        "Content Type": "Opinion Column",
        "Article Title": "Reimagining the permanent government: Reforming India's civil service architecture",
        "Post Content": "The permanent government and civil bureaucracy must transition from administrative gatekeepers to mission-driven enablers. Strengthening institutional accountability and regulatory capacity is critical for civic governance.",
        "URL": "https://swarajyamag.com/governance/reimagining-the-permanent-government-reforming-indias-civil-service-architecture/",
        "Primary Classification": "Public Policy & Governance",
        "Sub-classification": "Civic Systems"
    },
    {
        "ID": "SM-018",
        "Date": "2023-08-23",
        "Publication": "Livemint",
        "Co-Authors": "Suresh Prabhu",
        "Content Type": "Opinion Column",
        "Article Title": "India has an opportunity to lead the fourth industrial revolution",
        "Post Content": "India stands at a pivotal juncture where technological rise, digital adoption, and frugal innovation can position our industries at the forefront of the Fourth Industrial Revolution.",
        "URL": "https://www.livemint.com/opinion/online-views/india-has-an-opportunity-to-lead-the-fourth-industrial-revolution-11692795493821.html",
        "Primary Classification": "Technology & Innovation",
        "Sub-classification": "Frugal Innovation"
    },
    {
        "ID": "SM-019",
        "Date": "2023-07-26",
        "Publication": "Livemint",
        "Co-Authors": "Suresh Prabhu",
        "Content Type": "Opinion Column",
        "Article Title": "Universities should lead the way as we adapt to the future of work",
        "Post Content": "As artificial intelligence and automation reshape industry demands, universities must overhaul curricula to prioritize critical thinking, adaptability, and multidisciplinary skills for the future of work.",
        "URL": "https://www.livemint.com/opinion/online-views/universities-should-lead-the-way-as-we-adapt-to-the-future-of-work-11690374182903.html",
        "Primary Classification": "Education & Pedagogy",
        "Sub-classification": "Higher Education"
    },
    {
        "ID": "SM-020",
        "Date": "2023-06-15",
        "Publication": "The Economic Times",
        "Co-Authors": "Solo",
        "Content Type": "Opinion Column",
        "Article Title": "On academic disciplines, India could take a leaf out of US' book",
        "Post Content": "Compartmentalizing education into rigid silos limits holistic thinking. Interdisciplinary learning and modern pedagogy are essential for equipping students with problem-solving skills for complex global challenges.",
        "URL": "https://economictimes.indiatimes.com/opinion/et-commentary/on-academic-disciplines-india-could-take-a-leaf-out-of-us-book/",
        "Primary Classification": "Education & Pedagogy",
        "Sub-classification": "Pedagogy"
    },
    {
        "ID": "SM-021",
        "Date": "2023-01-04",
        "Publication": "Livemint",
        "Co-Authors": "Suresh Prabhu",
        "Content Type": "Opinion Column",
        "Article Title": "We need biographies written of our great academic institutions",
        "Post Content": "Academic institutions carry deep institutional memory and cultural values. Documenting their biographies and historical milestones inspires upcoming generations and reinforces institutional excellence.",
        "URL": "https://www.livemint.com/opinion/online-views/we-need-biographies-written-of-our-great-academic-institutions-11672837492819.html",
        "Primary Classification": "Education & Pedagogy",
        "Sub-classification": "Higher Education"
    },
    {
        "ID": "SM-022",
        "Date": "2022-12-01",
        "Publication": "Livemint",
        "Co-Authors": "Suresh Prabhu",
        "Content Type": "Opinion Column",
        "Article Title": "Indian universities must prepare for foreign competition at home",
        "Post Content": "The entry of foreign universities into India presents both opportunities and challenges. Domestic higher education institutions must elevate research quality and pedagogical rigor to thrive in an open ecosystem.",
        "URL": "https://www.livemint.com/opinion/online-views/indian-universities-must-prepare-for-foreign-competition-at-home-11669894729183.html",
        "Primary Classification": "Education & Pedagogy",
        "Sub-classification": "Higher Education"
    },
    {
        "ID": "SM-023",
        "Date": "2022-09-12",
        "Publication": "Swarajya",
        "Co-Authors": "Solo",
        "Content Type": "Opinion Column",
        "Article Title": "State capacity and civil society: Redefining NGO-government partnerships",
        "Post Content": "Civil society organizations and NGOs play an indispensable role in community-level governance and bridging policy delivery gaps. A collaborative model between the state and social innovators strengthens civic systems.",
        "URL": "https://swarajyamag.com/ideas/state-capacity-and-civil-society-redefining-ngo-government-partnerships",
        "Primary Classification": "Public Policy & Governance",
        "Sub-classification": "Civic Systems"
    },
    {
        "ID": "SM-024",
        "Date": "2020-03-09",
        "Publication": "Swarajya",
        "Co-Authors": "Solo",
        "Content Type": "Interview",
        "Article Title": "The Idea Of Rashtram Stems From The Urgent Need To Generate Advanced Leadership In India",
        "Post Content": "India requires visionary leaders grounded in the civilizational ethos of Bharat and capable of navigating modern public policy complexities. Rashtram aims to institutionalize nation-building leadership education.",
        "URL": "https://swarajyamag.com/interviews/the-idea-of-rashtram-stems-from-the-urgent-need-to-generate-advanced-leadership-in-india-shobhit-mathur",
        "Primary Classification": "Education & Pedagogy",
        "Sub-classification": "Higher Education"
    },
    {
        "ID": "SM-025",
        "Date": "2016-05-18",
        "Publication": "The Hindu",
        "Co-Authors": "Nomesh Bolia",
        "Content Type": "Research Essay",
        "Article Title": "The MGNREGA index: Assessing implementation across states",
        "Post Content": "Empirical evaluation of rural employment schemes reveals critical disparities in wage payment delays and project completion. Developing robust civic indices is essential to evaluate state delivery mechanisms.",
        "URL": "https://www.thehindu.com/opinion/op-ed/the-mgnrega-index/article8612015.ece",
        "Primary Classification": "Public Policy & Governance",
        "Sub-classification": "Regulatory Frameworks"
    },
    {
        "ID": "SM-026",
        "Date": "2016-01-21",
        "Publication": "Swarajya",
        "Co-Authors": "Solo",
        "Content Type": "Opinion Column",
        "Article Title": "More NGOs In India Need To Work With The Government Instead Of Opposing It",
        "Post Content": "Constructive civic development occurs when non-governmental organizations partner directly with state administrative machinery rather than engaging in reflexive antagonism.",
        "URL": "https://swarajyamag.com/politics/more-ngos-in-india-need-to-work-with-the-government-instead-of-opposing-it",
        "Primary Classification": "Public Policy & Governance",
        "Sub-classification": "Civic Systems"
    },
    {
        "ID": "SM-027",
        "Date": "2024-08-20",
        "Publication": "Financial Express",
        "Co-Authors": "Solo",
        "Content Type": "Opinion Column",
        "Article Title": "Integrating Indian Knowledge Systems with NEP 2020: Building Leadership for Gen Z",
        "Post Content": "Integrating Indian Knowledge Systems into modern university curricula enables students to develop holistic leadership, cultural grounding, and multidimensional problem-solving.",
        "URL": "https://www.financialexpress.com/jobs-career/education/integrating-indian-knowledge-systems-with-nep-2020-building-leadership-for-gen-z/3588921/",
        "Primary Classification": "Education & Pedagogy",
        "Sub-classification": "Higher Education"
    },
    {
        "ID": "SM-028",
        "Date": "2018-02-11",
        "Publication": "IndiaSpend",
        "Co-Authors": "Pallabi Mandal",
        "Content Type": "Research Essay",
        "Article Title": "Haphazard Appointment Of Judges Worsening High Court Judicial Delays",
        "Post Content": "Empirical analysis of Department of Justice data shows ad-hoc judge allocations exacerbate pendency. Systematic judicial appointments and caseload rationalization are urgent for civic justice delivery.",
        "URL": "https://www.indiaspend.com/haphazard-appointment-of-judges-worsening-high-court-judicial-delays-83492/",
        "Primary Classification": "Public Policy & Governance",
        "Sub-classification": "Civic Systems"
    },
    {
        "ID": "SM-029",
        "Date": "2016-04-29",
        "Publication": "The Wire",
        "Co-Authors": "Nomesh Bolia",
        "Content Type": "Research Essay",
        "Article Title": "Worst Year Ever For MGNREGA, 23% Drop In Assets Created",
        "Post Content": "Assessing rural welfare schemes reveals sharp contractions in physical asset creation and wage disbursements. Regulatory efficiency and systemic accountability must guide rural social programs.",
        "URL": "https://thewire.in/economy/worst-year-ever-for-mgnrega-23-drop-in-assets-created",
        "Primary Classification": "Public Policy & Governance",
        "Sub-classification": "Regulatory Frameworks"
    },
    {
        "ID": "SM-030",
        "Date": "2025-02-03",
        "Publication": "NDTV",
        "Co-Authors": "Suresh Prabhu",
        "Content Type": "Opinion Column",
        "Article Title": "Outdated Laws Are Stifling India's Private Universities",
        "Post Content": "Regulatory rigidity and outdated governance laws are stifling private universities in India, preventing academic innovation, flexible student intake, and multidisciplinary education models essential for global competitiveness.",
        "URL": "https://www.ndtv.com/opinion/outdated-laws-are-stifling-indias-private-universities-7860124",
        "Primary Classification": "Education & Pedagogy",
        "Sub-classification": "Systemic Reform"
    },
    {
        "ID": "SM-031",
        "Date": "2017-03-31",
        "Publication": "The Quint",
        "Co-Authors": "Solo",
        "Content Type": "Opinion Column",
        "Article Title": "Why Judicial Reform Should Start With Lower Courts",
        "Post Content": "Immediate improvements in the judicial system could be achieved by focusing on lower courts, collecting standardized case data, and horizontally replicating procedural innovations.",
        "URL": "https://www.thequint.com/voices/opinion/indian-judicial-system-needs-a-reform-which-starts-at-lower-courts",
        "Primary Classification": "Public Policy & Governance",
        "Sub-classification": "Civic Systems"
    },
    {
        "ID": "SM-032",
        "Date": "2026-06-15",
        "Publication": "LinkedIn Pulse / Print",
        "Co-Authors": "Suresh Prabhu",
        "Content Type": "Opinion Column",
        "Article Title": "Why education must teach systems thinking",
        "Post Content": "For over a century, modern education has been organised around disciplines. The defining challenges of our time emerge at intersections. The student trained exclusively within a single discipline is increasingly unprepared for the reality they will inherit. What is missing is ways of thinking that enable individuals to operate in a world of deeply interconnected systems.",
        "URL": "https://www.linkedin.com/in/shobhitmathur/recent-activity/shares/",
        "Primary Classification": "Education & Pedagogy",
        "Sub-classification": "Pedagogy"
    }
]

def generate_master_workbook():
    master_xlsx = "shobhit_mathur_master_sheet.xlsx"
    df = pd.DataFrame(MASTER_DATA)
    
    wb = openpyxl.Workbook()
    
    # Styles
    navy_fill = PatternFill(start_color="0F172A", end_color="0F172A", fill_type="solid") # Slate 900
    accent_fill = PatternFill(start_color="1E40AF", end_color="1E40AF", fill_type="solid") # Royal Blue
    edu_fill = PatternFill(start_color="047857", end_color="047857", fill_type="solid") # Emerald Green
    policy_fill = PatternFill(start_color="B45309", end_color="B45309", fill_type="solid") # Amber 700
    tech_fill = PatternFill(start_color="6D28D9", end_color="6D28D9", fill_type="solid") # Purple 700
    
    header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
    title_font = Font(name="Segoe UI", size=14, bold=True, color="0F172A")
    data_font = Font(name="Segoe UI", size=10)
    bold_data = Font(name="Segoe UI", size=10, bold=True)
    link_font = Font(name="Segoe UI", size=10, color="2563EB", underline="single")
    
    thin_border = Border(
        left=Side(style='thin', color='E2E8F0'),
        right=Side(style='thin', color='E2E8F0'),
        top=Side(style='thin', color='E2E8F0'),
        bottom=Side(style='thin', color='E2E8F0')
    )
    zebra_fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
    
    headers = [
        "ID", "Date", "Publication", "Co-Authors", "Content Type", 
        "Article Title", "Post Content (Author Commentary)", "URL", 
        "Primary Classification", "Sub-classification"
    ]
    
    # Helper to style a table sheet
    def build_table_sheet(ws, dataset, tab_fill):
        ws.views.sheetView[0].showGridLines = True
        ws.append(headers)
        ws.row_dimensions[1].height = 28
        
        for col_idx in range(1, len(headers) + 1):
            cell = ws.cell(row=1, column=col_idx)
            cell.fill = tab_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")
            
        for row_idx, item in enumerate(dataset, start=2):
            row_vals = [
                item["ID"], item["Date"], item["Publication"], item["Co-Authors"], item["Content Type"],
                item["Article Title"], item["Post Content"], item["URL"],
                item["Primary Classification"], item["Sub-classification"]
            ]
            ws.append(row_vals)
            ws.row_dimensions[row_idx].height = 65
            is_zebra = (row_idx % 2 == 0)
            
            for col_idx in range(1, len(headers) + 1):
                cell = ws.cell(row=row_idx, column=col_idx)
                cell.border = thin_border
                if is_zebra:
                    cell.fill = zebra_fill
                
                col_name = headers[col_idx - 1]
                if col_name in ["ID", "Date", "Publication", "Co-Authors", "Content Type", "Primary Classification", "Sub-classification"]:
                    cell.alignment = Alignment(horizontal="center", vertical="top")
                    cell.font = data_font
                elif col_name in ["Article Title", "Post Content (Author Commentary)"]:
                    cell.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
                    cell.font = data_font
                elif col_name == "URL":
                    cell.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
                    cell.font = link_font
                    import urllib.parse
                    title = row_vals[headers.index("Article Title")]
                    pub = row_vals[headers.index("Publication")]
                    query = urllib.parse.quote(f'"{title}" {pub} Shobhit Mathur')
                    search_url = f'https://www.google.com/search?q={query}'
                    cell.value = f'=HYPERLINK("{search_url}", "Search on Google")'
                    
        widths = {
            "ID": 10, "Date": 13, "Publication": 20, "Co-Authors": 18, "Content Type": 18,
            "Article Title": 38, "Post Content (Author Commentary)": 50, "URL": 38,
            "Primary Classification": 24, "Sub-classification": 24
        }
        for col_idx, col_name in enumerate(headers, start=1):
            ws.column_dimensions[get_column_letter(col_idx)].width = widths[col_name]

    # 1. Master Table Sheet
    ws_master = wb.active
    ws_master.title = "Master Database"
    build_table_sheet(ws_master, MASTER_DATA, navy_fill)
    
    # 2. Category Tabs
    edu_data = [d for d in MASTER_DATA if d["Primary Classification"] == "Education & Pedagogy"]
    ws_edu = wb.create_sheet(title="Education & Pedagogy")
    build_table_sheet(ws_edu, edu_data, edu_fill)
    
    policy_data = [d for d in MASTER_DATA if d["Primary Classification"] == "Public Policy & Governance"]
    ws_policy = wb.create_sheet(title="Public Policy & Governance")
    build_table_sheet(ws_policy, policy_data, policy_fill)
    
    tech_data = [d for d in MASTER_DATA if d["Primary Classification"] == "Technology & Innovation"]
    ws_tech = wb.create_sheet(title="Technology & Innovation")
    build_table_sheet(ws_tech, tech_data, tech_fill)
    
    # 3. Executive Dashboard & Metrics Sheet
    ws_dash = wb.create_sheet(title="Executive Dashboard", index=0)
    ws_dash.views.sheetView[0].showGridLines = True
    
    ws_dash.cell(row=2, column=2, value="Shobhit Mathur — Executive Content Intelligence Dashboard").font = Font(name="Segoe UI", size=16, bold=True, color="0F172A")
    ws_dash.cell(row=3, column=2, value=f"Comprehensive Master Repository of Extracted Articles, Publications, and Feed Commentary | Total Records: {len(df)}").font = Font(name="Segoe UI", size=11, italic=True, color="64748B")
    
    # Metrics Cards
    metrics = [
        ("Total Articles", len(df), "0F172A"),
        ("Education & Pedagogy", len(edu_data), "047857"),
        ("Public Policy & Gov", len(policy_data), "B45309"),
        ("Tech & Innovation", len(tech_data), "6D28D9")
    ]
    
    for m_idx, (label, val, color_code) in enumerate(metrics):
        c_col = 2 + (m_idx * 2) # B=2, D=4, F=6, H=8
        # Value box
        ws_dash.cell(row=5, column=c_col, value=val).font = Font(name="Segoe UI", size=20, bold=True, color=color_code)
        ws_dash.cell(row=5, column=c_col).alignment = Alignment(horizontal="center", vertical="center")
        ws_dash.cell(row=5, column=c_col).fill = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")
        ws_dash.cell(row=5, column=c_col).border = thin_border
        
        # Label box
        ws_dash.cell(row=6, column=c_col, value=label).font = Font(name="Segoe UI", size=9, bold=True, color="475569")
        ws_dash.cell(row=6, column=c_col).alignment = Alignment(horizontal="center", vertical="center")
        ws_dash.cell(row=6, column=c_col).fill = PatternFill(start_color="E2E8F0", end_color="E2E8F0", fill_type="solid")
        ws_dash.cell(row=6, column=c_col).border = thin_border
        
    ws_dash.row_dimensions[5].height = 36
    ws_dash.row_dimensions[6].height = 20
    
    # Table 1: Publication Breakdown
    ws_dash.merge_cells("B9:D9")
    ws_dash.cell(row=9, column=2, value="Breakdown by Publication Source").font = Font(name="Segoe UI", size=12, bold=True, color="0F172A")
    
    ws_dash.merge_cells("B10:C10")
    ws_dash.cell(row=10, column=2, value="Publication").font = header_font
    ws_dash.cell(row=10, column=2).fill = navy_fill
    ws_dash.cell(row=10, column=3).fill = navy_fill # Fill merged part
    ws_dash.cell(row=10, column=4, value="Count").font = header_font
    ws_dash.cell(row=10, column=4).fill = navy_fill
    ws_dash.cell(row=10, column=4).alignment = Alignment(horizontal="center")
    
    pub_counts = df["Publication"].value_counts().reset_index()
    pub_counts.columns = ["Publication", "Count"]
    for p_idx, p_row in pub_counts.iterrows():
        r = 11 + p_idx
        ws_dash.merge_cells(start_row=r, start_column=2, end_row=r, end_column=3)
        ws_dash.cell(row=r, column=2, value=p_row["Publication"]).font = data_font
        ws_dash.cell(row=r, column=2).border = thin_border
        ws_dash.cell(row=r, column=3).border = thin_border
        ws_dash.cell(row=r, column=4, value=p_row["Count"]).font = bold_data
        ws_dash.cell(row=r, column=4).border = thin_border
        ws_dash.cell(row=r, column=4).alignment = Alignment(horizontal="center")
        
    # Table 2: Sub-Classification Breakdown
    ws_dash.merge_cells("F9:H9")
    ws_dash.cell(row=9, column=6, value="Taxonomy Sub-Classifications").font = Font(name="Segoe UI", size=12, bold=True, color="0F172A")
    
    ws_dash.merge_cells("F10:G10")
    ws_dash.cell(row=10, column=6, value="Sub-classification").font = header_font
    ws_dash.cell(row=10, column=6).fill = accent_fill
    ws_dash.cell(row=10, column=7).fill = accent_fill # Fill merged part
    ws_dash.cell(row=10, column=8, value="Count").font = header_font
    ws_dash.cell(row=10, column=8).fill = accent_fill
    ws_dash.cell(row=10, column=8).alignment = Alignment(horizontal="center")
    
    sub_counts = df["Sub-classification"].value_counts().reset_index()
    sub_counts.columns = ["Sub-classification", "Count"]
    for s_idx, s_row in sub_counts.iterrows():
        r = 11 + s_idx
        ws_dash.merge_cells(start_row=r, start_column=6, end_row=r, end_column=7)
        ws_dash.cell(row=r, column=6, value=s_row["Sub-classification"]).font = data_font
        ws_dash.cell(row=r, column=6).border = thin_border
        ws_dash.cell(row=r, column=7).border = thin_border
        ws_dash.cell(row=r, column=8, value=s_row["Count"]).font = bold_data
        ws_dash.cell(row=r, column=8).border = thin_border
        ws_dash.cell(row=r, column=8).alignment = Alignment(horizontal="center")
        
    ws_dash.column_dimensions['B'].width = 22
    ws_dash.column_dimensions['C'].width = 4
    ws_dash.column_dimensions['D'].width = 22
    ws_dash.column_dimensions['E'].width = 6
    ws_dash.column_dimensions['F'].width = 22
    ws_dash.column_dimensions['G'].width = 4
    ws_dash.column_dimensions['H'].width = 22
    
    wb.save(master_xlsx)
    print(f"Successfully generated Master Sheet: {master_xlsx}")

if __name__ == "__main__":
    generate_master_workbook()
